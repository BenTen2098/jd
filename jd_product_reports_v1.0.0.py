from concurrent.futures import ThreadPoolExecutor, as_completed
from datetime import date, datetime, timedelta
import json
from pathlib import Path
import tempfile
import time
from urllib.parse import quote
import zipfile

import openpyxl
import requests


MANAGE_CREATE_URL = "https://jdsz.jd.com/brand/reportCenter/recommendReport/downloadManageProData.ajax"
MANAGE_LIST_URL = "https://jdsz.jd.com/brand/reportCenter/myReport/getReportList.ajax"
MANAGE_DOWNLOAD_URL = "https://jdsz.jd.com/brand/reportCenter/myReport/downLoadReport.ajax"
SUPPLY_CREATE_URL = "https://zhgateway.jd.com/inventoryajax/reportCenter/recommendReport/downloadERPSupplyChainProData.ajax"
SUPPLY_LIST_URL = MANAGE_LIST_URL

MANAGE_CREATE_UUID = "869c31e7fce14d03a5b0-1a0d12c56f1"
MANAGE_LIST_UUID = "64eba73e0c092db8444e-1a0d13a535d"
MANAGE_DOWNLOAD_UUID = "1d80b1889dd8c9e0bc08-1a0d148547b"
SUPPLY_CREATE_UUID = "1b8f7d2d-bff7-4fec-9551-3106215053b6"
SUPPLY_LIST_UUID = "b79d06283f967916b639-1a0d1551d29"

MANAGE_CREATE_MNP = "68326260128cee361dd275dc1baa4bc6"
MANAGE_LIST_MNP = "685d4acb875a26145bed705a6933c5af"
MANAGE_DOWNLOAD_MNP = "a22e9561c3cfdd5831b903b389dba5b1"
SUPPLY_CREATE_MNP = "6756e958093acfb8b88b5a908d9d1d75"
SUPPLY_LIST_MNP = "888a94de2c18b79c914f2d928b8bfa07"

USER_AGENT = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/151.0.0.0 Safari/537.36"


def api_headers(cookie, p_pin, user_mnp, referer, content_type=None):
    headers = {
        "Accept": "application/json, text/plain, */*",
        "Accept-Language": "zh-CN,zh-TW;q=0.9,zh;q=0.8,en-US;q=0.7,en;q=0.6",
        "Cookie": cookie,
        "Origin": "https://jdsz.jd.com",
        "p-pin": quote(p_pin, safe=""),
        "Referer": referer,
        "User-Agent": USER_AGENT,
        "user-mnp": user_mnp,
        "user-mup": str(int(time.time() * 1000)),
        "X-Requested-With": "XMLHttpRequest",
    }
    if content_type:
        headers["Content-Type"] = content_type
    return headers


def build_manage_payload(business_date):
    end_date = datetime.strptime(business_date, "%Y-%m-%d").date()
    start_date = end_date - timedelta(days=30)
    payload = {
        "startDate": start_date.isoformat(),
        "endDate": business_date,
        "date": f"10{business_date}",
        "brandId": "all",
        "thirdCategoryId": "all",
        "provinceId": "all",
        "cityId": "all",
        "formFlag": "0",
        "projectId": "100146",
        "secondSourceId": "2002",
        "skuId": "",
        "firstBuTypeId": "",
        "secondBuTypeId": "",
        "isRdc": "0",
        "saleChannal": "1",
        "proType": "SKU",
        "dateType": "",
    }
    query_search = dict(payload)
    query_search["dateType"] = "dayRange"
    query_search["categoryReset"] = {}
    query_search["provinceReset"] = {}
    payload["querySearch"] = json.dumps(query_search, ensure_ascii=False, separators=(",", ":"))
    return payload


def build_supply_dates(business_date):
    end_date = datetime.strptime(business_date, "%Y-%m-%d").date()
    return [(end_date - timedelta(days=offset)).isoformat() for offset in range(7)]


def build_supply_payload(business_date):
    return {
        "isRdc": "0",
        "brandId": "all",
        "firstCategoryId": "",
        "secondCategoryId": "",
        "thirdCategoryId": "all",
        "date": business_date,
        "startDate": business_date,
        "endDate": business_date,
        "skuId": "",
        "skuStatusCd": "",
        "dataType": "offline",
        "id": 1,
        "excludeEmpty": "0",
    }


def build_output_filename(business_date, table_name, timestamp_ms):
    return f"{timestamp_ms}-{business_date}-{table_name}.xlsx"


def validate_xlsx_file(path):
    path = Path(path)
    if not path.is_file() or path.stat().st_size == 0:
        raise RuntimeError(f"XLSX output is empty: {path}")
    try:
        with zipfile.ZipFile(path) as archive:
            if archive.testzip() is not None:
                raise RuntimeError(f"XLSX output is corrupt: {path}")
            names = set(archive.namelist())
    except zipfile.BadZipFile as exc:
        raise RuntimeError(f"XLSX output is not a ZIP workbook: {path}") from exc
    if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
        raise RuntimeError(f"XLSX workbook members are incomplete: {path}")


def download_binary(download_session, download_url, target_path, expected_suffix, magic):
    target_path = Path(target_path)
    part_path = target_path.with_suffix(target_path.suffix + ".part")
    part_path.unlink(missing_ok=True)
    written = 0
    with download_session.get(download_url, timeout=(10, 180), stream=True) as response:
        response.raise_for_status()
        content_type = response.headers.get("Content-Type", "").split(";", 1)[0].lower()
        expected_size = int(response.headers.get("Content-Length") or 0)
        with part_path.open("wb") as file_handle:
            for chunk in response.iter_content(chunk_size=1024 * 1024):
                if chunk:
                    file_handle.write(chunk)
                    written += len(chunk)
    allowed_types = {
        ".zip": {"", "application/zip", "application/octet-stream"},
        ".xlsx": {"", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/octet-stream"},
    }
    try:
        if expected_suffix not in allowed_types:
            raise RuntimeError(f"unsupported download suffix: {expected_suffix}")
        if content_type not in allowed_types[expected_suffix]:
            raise RuntimeError(f"unexpected download content type: {content_type}")
        if written == 0 or (expected_size and written != expected_size):
            raise RuntimeError(f"download size mismatch: expected={expected_size}, actual={written}")
        if part_path.read_bytes()[: len(magic)] != magic:
            raise RuntimeError(f"download magic does not match {expected_suffix}")
        part_path.replace(target_path)
    except Exception:
        part_path.unlink(missing_ok=True)
        raise
    return target_path


def append_store_name_column(source_path, target_path, store_name):
    workbook = openpyxl.load_workbook(source_path)
    for worksheet in workbook.worksheets:
        header_row = None
        for row_index in range(1, worksheet.max_row + 1):
            values = [worksheet.cell(row_index, column).value for column in range(1, worksheet.max_column + 1)]
            if sum(value is not None for value in values) >= 2:
                header_row = row_index
                break
        if header_row is None:
            continue
        header_columns = [
            cell.column for cell in worksheet[header_row] if cell.value is not None
        ]
        store_column = (max(header_columns) if header_columns else worksheet.max_column) + 1
        worksheet.cell(header_row, store_column).value = "店铺名"
        for row_index in range(header_row + 1, worksheet.max_row + 1):
            has_data = any(
                worksheet.cell(row_index, column).value is not None
                for column in range(1, store_column)
            )
            if has_data:
                worksheet.cell(row_index, store_column).value = store_name
    workbook.save(target_path)
    workbook.close()
    validate_xlsx_file(target_path)
    return target_path


def merge_supply_workbooks(source_paths, target_path):
    if not source_paths:
        raise RuntimeError("no supply-chain workbooks to merge")
    merged_workbook = openpyxl.load_workbook(source_paths[0])
    for source_path in source_paths[1:]:
        source_workbook = openpyxl.load_workbook(source_path, read_only=True, data_only=False)
        for source_sheet in source_workbook.worksheets:
            if source_sheet.title not in merged_workbook.sheetnames:
                target_sheet = merged_workbook.create_sheet(source_sheet.title)
                for row in source_sheet.iter_rows(values_only=True):
                    target_sheet.append(row)
                continue
            target_sheet = merged_workbook[source_sheet.title]
            header_row = 1
            for row_index, row in enumerate(source_sheet.iter_rows(values_only=True), start=1):
                if sum(value is not None for value in row) >= 2:
                    header_row = row_index
                    break
            for row in source_sheet.iter_rows(min_row=header_row + 1, values_only=True):
                if any(value is not None for value in row):
                    target_sheet.append(row)
        source_workbook.close()
    merged_workbook.save(target_path)
    merged_workbook.close()
    validate_xlsx_file(target_path)
    return target_path


def poll_report(api_session, cookie, p_pin, list_url, list_uuid, list_mnp, list_referer, report_type, report_marker, submitted_at, poll_interval_seconds, max_wait_seconds):
    deadline = time.monotonic() + max_wait_seconds
    selected_report = None
    while time.monotonic() < deadline:
        poll_result = None
        for attempt in range(1, 4):
            try:
                response = api_session.get(
                    list_url,
                    params={"uuid": list_uuid},
                    headers=api_headers(cookie, p_pin, list_mnp, list_referer),
                    timeout=(10, 60),
                )
                response.raise_for_status()
                if "json" not in response.headers.get("Content-Type", "").lower():
                    raise RuntimeError("report list did not return JSON")
                poll_result = response.json()
                break
            except (requests.ConnectionError, requests.Timeout):
                if attempt == 3:
                    raise
                time.sleep(attempt)
        if not isinstance(poll_result, dict) or poll_result.get("message") != "success":
            raise RuntimeError(f"report polling failed: {poll_result}")
        matching_reports = []
        for report in poll_result.get("content", {}).get("data", []):
            if report.get("reportType") != report_type or report_marker not in (report.get("reportName") or ""):
                continue
            created_at_text = report.get("createTime")
            try:
                created_at = datetime.strptime(created_at_text, "%Y-%m-%d %H:%M:%S")
            except (TypeError, ValueError):
                continue
            if created_at >= submitted_at - timedelta(seconds=30):
                matching_reports.append(report)
        if matching_reports:
            matching_reports.sort(key=lambda item: item.get("createTime", ""), reverse=True)
            selected_report = matching_reports[0]
            status = str(selected_report.get("status"))
            if status == "2":
                return selected_report
            if status not in {"1", "2"}:
                raise RuntimeError(f"report entered unexpected status: {selected_report}")
        time.sleep(poll_interval_seconds)
    raise TimeoutError(f"report was not ready within {max_wait_seconds} seconds: {report_marker}")


def create_supply_report(current_date, p_pin, cookie):
    payload = build_supply_payload(current_date)
    submitted_at = datetime.now()
    api_session = requests.Session()
    try:
        response = api_session.post(
            SUPPLY_CREATE_URL,
            params={"uuid": SUPPLY_CREATE_UUID},
            headers=api_headers(cookie, p_pin, SUPPLY_CREATE_MNP, "https://jdsz.jd.com/scbrandweb/brand/view/supplyReport/supplyChainPro.html", "application/json;charset=UTF-8"),
            json=payload,
            timeout=(10, 60),
        )
        response.raise_for_status()
        if "json" not in response.headers.get("Content-Type", "").lower():
            raise RuntimeError("supply-chain report creation did not return JSON")
        create_result = response.json()
        if str(create_result.get("status")) != "0":
            raise RuntimeError(f"supply-chain report creation failed: {create_result}")
        task_name = create_result.get("content", {}).get("taskName")
        if not task_name:
            raise RuntimeError("supply-chain report creation returned no taskName")
        return {"date": current_date, "submitted_at": submitted_at, "task_name": task_name}
    finally:
        api_session.close()


def run_report(report_kind, business_date, p_pin, cookie, output_dir, poll_interval_seconds=10, max_wait_seconds=1800):
    if report_kind not in {"manage_pro", "supply_chain_pro"}:
        raise ValueError("report_kind must be manage_pro or supply_chain_pro")
    datetime.strptime(business_date, "%Y-%m-%d")
    if not p_pin or not cookie:
        raise ValueError("p_pin and cookie are required")
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    api_session = requests.Session()
    download_session = requests.Session()
    timestamp_ms = int(time.time() * 1000)

    if report_kind == "manage_pro":
        submitted_at = datetime.now()
        payload = build_manage_payload(business_date)
        response = api_session.post(
            MANAGE_CREATE_URL,
            params={"uuid": MANAGE_CREATE_UUID},
            headers=api_headers(cookie, p_pin, MANAGE_CREATE_MNP, "https://jdsz.jd.com/brand/reportCenter/downLoadReport.html", "application/x-www-form-urlencoded"),
            data=payload,
            timeout=(10, 60),
        )
        response.raise_for_status()
        selected_report = poll_report(
            api_session, cookie, p_pin, MANAGE_LIST_URL, MANAGE_LIST_UUID, MANAGE_LIST_MNP,
            "https://jdsz.jd.com/brand/reportCenter/myReport.html", "managePro",
            f"{payload['startDate']}至{business_date}", submitted_at, poll_interval_seconds, max_wait_seconds,
        )
        response = api_session.get(
            MANAGE_DOWNLOAD_URL,
            params={"id": selected_report.get("reportId"), "uuid": MANAGE_DOWNLOAD_UUID},
            headers=api_headers(cookie, p_pin, MANAGE_DOWNLOAD_MNP, "https://jdsz.jd.com/brand/reportCenter/myReport.html"),
            timeout=(10, 60),
        )
        response.raise_for_status()
        if "json" not in response.headers.get("Content-Type", "").lower():
            raise RuntimeError("manage report download endpoint did not return JSON")
        download_url = response.json().get("content", {}).get("result")
        if not download_url:
            raise RuntimeError("manage report download response contains no result URL")
        output_path = output_dir / build_output_filename(business_date, "经营状况-商品明细报表", timestamp_ms)
        final_part = output_path.with_suffix(output_path.suffix + ".part")
        with tempfile.TemporaryDirectory(prefix="manage_report_") as temp_dir:
            zip_path = download_binary(download_session, download_url, Path(temp_dir) / "manage_report.zip", ".zip", b"PK")
            with zipfile.ZipFile(zip_path) as archive:
                xlsx_members = [name for name in archive.namelist() if name.lower().endswith(".xlsx")]
                if not xlsx_members:
                    raise RuntimeError("manage report ZIP contains no XLSX")
                raw_xlsx = Path(temp_dir) / "manage_report.xlsx"
                with archive.open(xlsx_members[0]) as source, raw_xlsx.open("wb") as target:
                    target.write(source.read())
            append_store_name_column(raw_xlsx, final_part, p_pin)
        final_part.replace(output_path)
        return output_path

    supply_dates = build_supply_dates(business_date)
    created_reports = {}
    with ThreadPoolExecutor(max_workers=len(supply_dates)) as executor:
        futures = {
            executor.submit(create_supply_report, current_date, p_pin, cookie): current_date
            for current_date in supply_dates
        }
        for future in as_completed(futures):
            created_report = future.result()
            created_reports[created_report["date"]] = created_report
    if set(created_reports) != set(supply_dates):
        raise RuntimeError("not all supply-chain report requests were created")

    table_name = "供应链库存-商品明细"
    with tempfile.TemporaryDirectory(prefix="supply_reports_") as temp_dir:
        completed_reports = {}
        for current_date in supply_dates:
            created_report = created_reports[current_date]
            completed_reports[current_date] = poll_report(
                api_session, cookie, p_pin, SUPPLY_LIST_URL, SUPPLY_LIST_UUID, SUPPLY_LIST_MNP,
                "https://jdsz.jd.com/brand/reportCenter/myReport.html", "supplyChainPro",
                created_report["task_name"], created_report["submitted_at"], poll_interval_seconds, max_wait_seconds,
            )

        supply_output_paths = []
        for current_date in supply_dates:
            selected_report = completed_reports[current_date]
            download_url = selected_report.get("downloadLink")
            if not download_url:
                raise RuntimeError(f"supply-chain report has no downloadLink: {current_date}")
            raw_path = download_binary(download_session, download_url, Path(temp_dir) / f"{current_date}.xlsx", ".xlsx", b"PK\x03\x04")
            daily_path = Path(temp_dir) / f"{current_date}-with-store.xlsx"
            append_store_name_column(raw_path, daily_path, p_pin)
            supply_output_paths.append(daily_path)
        output_path = output_dir / build_output_filename(business_date, table_name, timestamp_ms)
        final_part = output_path.with_suffix(output_path.suffix + ".part")
        merge_supply_workbooks(supply_output_paths, final_part)
    final_part.replace(output_path)
    return output_path


if __name__ == "__main__":
    business_date = "2026-09-23"
    p_pin = "官栈自营企业号"
    cookie = ""
    output_dir = Path(__file__).resolve().parent / "output"
    poll_interval_seconds = 10
    max_wait_seconds = 1800

    if not cookie:
        raise SystemExit("请在 __main__ 中填入已登录京东页面复制的 cookie")
    print(run_report("manage_pro", business_date, p_pin, cookie, output_dir, poll_interval_seconds, max_wait_seconds))
    print(run_report("supply_chain_pro", business_date, p_pin, cookie, output_dir, poll_interval_seconds, max_wait_seconds))
